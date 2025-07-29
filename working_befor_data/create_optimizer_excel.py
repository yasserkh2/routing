import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import re
import os

def parse_optimizer_results(file_path):
    """Parse the optimizer test results file and extract data for each profile and scenario."""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split the content by profile sections
    profile_sections = {}
    current_profile = None
    current_scenario = None
    
    # Regular expressions to match profile and scenario headers
    profile_scenario_pattern = re.compile(r'--- (\d+)\. Profile: ([^-]+) - ([^-]+) ---')
    
    # Split the content by lines
    lines = content.strip().split('\n')
    
    for i, line in enumerate(lines):
        # Check if this line is a profile/scenario header
        match = profile_scenario_pattern.match(line)
        if match:
            scenario_num, profile_name, scenario_name = match.groups()
            profile_name = profile_name.strip()
            scenario_name = scenario_name.strip()
            
            if profile_name not in profile_sections:
                profile_sections[profile_name] = {}
            
            current_profile = profile_name
            current_scenario = scenario_name
            profile_sections[current_profile][current_scenario] = {
                'headers': [],
                'data': []
            }
            
            # The next line should be the table header
            if i + 1 < len(lines) and '|' in lines[i + 1]:
                headers = [h.strip() for h in lines[i + 1].split('|')[1:-1]]
                profile_sections[current_profile][current_scenario]['headers'] = headers
            
            # Skip the header and separator lines
            i += 3
            
            # Extract data rows until we hit the next section or end of file
            while i < len(lines) and lines[i].startswith('|') and not lines[i].startswith('|--'):
                row_data = [cell.strip() for cell in lines[i].split('|')[1:-1]]
                profile_sections[current_profile][current_scenario]['data'].append(row_data)
                i += 1
    
    return profile_sections

def create_excel_file(profile_sections, output_path):
    """Create an Excel file with the parsed optimizer results."""
    wb = openpyxl.Workbook()
    
    # Remove the default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    # Define styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    data_alignment = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Create a consolidated sheet for each profile
    for profile_name, scenarios in profile_sections.items():
        # Create a new sheet for this profile
        ws = wb.create_sheet(title=profile_name[:31])  # Excel sheet names are limited to 31 characters
        
        # Add headers with Round Type column
        headers = ['Round Type'] + scenarios[list(scenarios.keys())[0]]['headers']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            
            # Set column width based on header length
            ws.column_dimensions[get_column_letter(col)].width = max(len(header) + 2, 15)
        
        # Start row for data
        row = 2
        
        # Process each scenario
        for scenario_name, scenario_data in scenarios.items():
            # Add data rows for this scenario
            for data_row in scenario_data['data']:
                # Add round type in first column
                cell = ws.cell(row=row, column=1, value=scenario_name)
                cell.alignment = data_alignment
                cell.border = thin_border
                
                # Add the rest of the data
                for col, value in enumerate(data_row, 2):  # Start from column 2 (after Round Type)
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.alignment = data_alignment
                    cell.border = thin_border
                    
                    # Format numeric columns
                    if col in [4, 5]:  # Traffic % and SLA % (shifted by 1 due to Round Type column)
                        try:
                            # Extract numeric value from percentage
                            numeric_value = float(value.strip('%'))
                            cell.value = numeric_value / 100  # Convert to decimal for Excel percentage format
                            cell.number_format = '0.0%'
                        except ValueError:
                            pass
                    elif col in [6, 7]:  # Price and Cost (shifted by 1 due to Round Type column)
                        try:
                            # Extract numeric value from currency format
                            numeric_value = float(value.replace('$', '').strip())
                            cell.value = numeric_value
                            cell.number_format = '$0.0000'
                        except ValueError:
                            pass
                
                row += 1
    
    # Save the workbook
    wb.save(output_path)
    print(f"Excel file created successfully: {output_path}")

def main():
    # Define file paths
    input_file = os.path.join(os.path.dirname(__file__), 'optimizer_test_results.txt')
    output_file = os.path.join(os.path.dirname(__file__), 'optimizer_test_results_new.xlsx')
    
    # Parse the optimizer results
    profile_sections = parse_optimizer_results(input_file)
    
    # Create the Excel file
    create_excel_file(profile_sections, output_file)

if __name__ == "__main__":
    main()